from paddleocr import PaddleOCR
from pdf2image import convert_from_path
from docx import Document
import os

# Initialize OCR
ocr = PaddleOCR(
    lang="en",
    use_doc_orientation_classify=False,
    use_doc_unwarping=False,
    use_textline_orientation=False,
    device="cpu"
)

file_path = input("Enter image or PDF path: ").strip()

if not os.path.exists(file_path):
    print("Error: File not found!")
    exit()

# Create output folder structure
OUTPUT_DIR = "output"
os.makedirs(OUTPUT_DIR, exist_ok=True)

base_name = os.path.splitext(os.path.basename(file_path))[0]

output_folder = os.path.join(
    OUTPUT_DIR,
    f"{base_name}_output"
)
os.makedirs(output_folder, exist_ok=True)

# Folder for temporary PDF page images
pages_folder = os.path.join(output_folder, "pages")
os.makedirs(pages_folder, exist_ok=True)

all_text = []

# ---------- PDF OCR ----------
if file_path.lower().endswith(".pdf"):

    print("\nProcessing PDF...\n")

    pages = convert_from_path(file_path)

    for page_num, page in enumerate(pages, start=1):

        image_name = os.path.join(
            pages_folder,
            f"page_{page_num}.png"
        )

        page.save(image_name, "PNG")

        result = ocr.predict(image_name)

        all_text.append(f"\n===== PAGE {page_num} =====\n")

        for page_result in result:
            for text in page_result["rec_texts"]:
                print(text)
                all_text.append(text)

        # Delete temporary image after OCR
        os.remove(image_name)

# ---------- IMAGE OCR ----------
else:

    print("\nProcessing Image...\n")

    result = ocr.predict(file_path)

    for page_result in result:
        for text in page_result["rec_texts"]:
            print(text)
            all_text.append(text)

# ---------- Save TXT ----------
txt_file = os.path.join(
    output_folder,
    f"{base_name}.txt"
)

with open(txt_file, "w", encoding="utf-8") as f:
    f.write("\n".join(all_text))

# ---------- Save DOCX ----------
doc = Document()

doc.add_heading(
    f"OCR Output - {base_name}",
    level=1
)

for line in all_text:
    doc.add_paragraph(line)

docx_file = os.path.join(
    output_folder,
    f"{base_name}.docx"
)

doc.save(docx_file)

# Remove pages folder if empty
try:
    os.rmdir(pages_folder)
except OSError:
    pass

print("\nOCR Completed Successfully!")
print(f"TXT saved to: {txt_file}")
print(f"DOCX saved to: {docx_file}")