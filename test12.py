"""
OCR with Table Detection & Reconstruction
==========================================

Requires:
    pip install paddleocr pdf2image opencv-python python-docx beautifulsoup4 lxml

paddleocr>=3.0 is required — this uses the PPStructureV3 pipeline (the
PaddleOCR 2.x `PPStructure` class was renamed/replaced by `PPStructureV3`
in the 3.x line; importing the old name raises ImportError on 3.x).
Poppler must be installed on the system for pdf2image to work.

What changed vs. the plain-OCR version:
  - Uses PaddleOCR's PPStructureV3 pipeline, which performs LAYOUT ANALYSIS
    first (splits a page into regions: title, text, table, figure, etc.),
    then runs TABLE STRUCTURE RECOGNITION on any region classified as table.
  - Each result's res["table_res_list"] holds one entry per detected table,
    each with a "pred_html" field. We parse that HTML and rebuild a *real*
    Word table (doc.add_table) with the correct row/column structure,
    instead of dumping cell text as flat paragraphs.
  - Non-table content comes from res["parsing_res_list"], which holds
    ordered blocks (text/title/etc.) with their recognized text — these
    are written as paragraphs, same as before.
  - "Correction" pass: cell text is stripped/whitespace-normalized, and
    rowspan/colspan are expanded into repeated cells so the grid stays
    rectangular.
"""

import os
import re
import time
import warnings

import cv2
import numpy as np
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt
from pdf2image import convert_from_path, pdfinfo_from_path
from paddleocr import PPStructureV3

# ==========================================================
# SETTINGS
# ==========================================================

warnings.filterwarnings("ignore")

os.environ["GLOG_minloglevel"] = "3"
os.environ["FLAGS_logtostderr"] = "0"
os.environ["FLAGS_use_mkldnn"] = "0"

OUTPUT_DIR = "output"

# ==========================================================
# IMAGE PREPROCESSING
# ==========================================================

def preprocess_image(image):
    """Light denoise only — PPStructure expects a normal BGR image,
    so we no longer collapse it to gray-as-BGR, just clean noise."""

    denoised = cv2.fastNlMeansDenoisingColored(
        image, None, 10, 10, 7, 21
    )
    return denoised


# ==========================================================
# STRUCTURE ENGINE (layout + table + OCR)
# ==========================================================

print("Loading layout/table/OCR model...")

structure_engine = PPStructureV3(
    use_doc_orientation_classify=True,
    use_doc_unwarping=True,
    use_textline_orientation=True,
    use_table_recognition=True,
    lang="en",
)

print("✓ Model loaded")


# ==========================================================
# TABLE HTML -> CLEAN GRID -> DOCX TABLE
# ==========================================================

def html_table_to_grid(html):
    """Parse PPStructure's HTML table output into a rectangular grid of
    cell strings, expanding rowspan/colspan so every row has equal columns."""

    soup = BeautifulSoup(html, "lxml")
    table_tag = soup.find("table")
    if table_tag is None:
        return []

    rows = table_tag.find_all("tr")
    grid = []
    # track cells already filled by a previous row's rowspan
    pending_rowspans = {}  # col_index -> (remaining_rows, text)

    for row in rows:
        cells = row.find_all(["td", "th"])
        current_row = []
        col_index = 0

        # first, fill in any rowspans carried over from above
        def fill_pending(col_index, current_row):
            while col_index in pending_rowspans:
                remaining, text = pending_rowspans[col_index]
                current_row.append(text)
                remaining -= 1
                if remaining <= 0:
                    del pending_rowspans[col_index]
                else:
                    pending_rowspans[col_index] = (remaining, text)
                col_index += 1
            return col_index

        col_index = fill_pending(col_index, current_row)

        for cell in cells:
            col_index = fill_pending(col_index, current_row)

            text = cell.get_text(strip=True)
            colspan = int(cell.get("colspan", 1))
            rowspan = int(cell.get("rowspan", 1))

            for _ in range(colspan):
                current_row.append(text)
                if rowspan > 1:
                    pending_rowspans[col_index] = (rowspan - 1, text)
                col_index += 1

        col_index = fill_pending(col_index, current_row)
        grid.append(current_row)

    # pad ragged rows so the grid is rectangular
    max_cols = max((len(r) for r in grid), default=0)
    for r in grid:
        while len(r) < max_cols:
            r.append("")

    return grid


def correct_grid(grid):
    """Clean up common OCR/table-structure artifacts:
       - strip stray whitespace / newlines inside cells
       - collapse repeated internal spaces
       - leave genuinely empty cells as '' (not dropped, to keep alignment)
    """
    cleaned = []
    for row in grid:
        new_row = []
        for cell in row:
            cell = re.sub(r"\s+", " ", cell or "").strip()
            new_row.append(cell)
        cleaned.append(new_row)
    return cleaned


def add_docx_table(doc, grid):
    if not grid or not grid[0]:
        return

    n_rows = len(grid)
    n_cols = len(grid[0])

    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.style = "Table Grid"

    for r, row_data in enumerate(grid):
        for c, cell_text in enumerate(row_data):
            cell = table.cell(r, c)
            cell.text = cell_text
            if r == 0:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True

    doc.add_paragraph("")  # spacing after table


# ==========================================================
# REGION PROCESSING (PPStructureV3 result schema)
# ==========================================================
#
# PPStructureV3's predict() returns a list of result objects (one per
# image/page). Each result behaves like a dict via result["res"] /
# result.json, containing (among others):
#
#   res["table_res_list"]   -> list of {"pred_html": "<table>...</table>", ...}
#   res["parsing_res_list"] -> ordered list of blocks, each typically like
#                              {"block_label": "text"/"title"/"table"/...,
#                               "block_content": "...", "block_bbox": [...]}
#
# Table blocks inside parsing_res_list are cross-referenced with
# table_res_list by order; we render the table once (from table_res_list)
# and skip re-printing its raw block_content as a paragraph.

def get_res_dict(result_item):
    """Normalize a PPStructureV3 result item into a plain dict."""
    if hasattr(result_item, "json"):
        data = result_item.json
        # some versions wrap actual payload one level deeper under "res"
        if isinstance(data, dict) and "res" in data:
            return data["res"]
        return data
    if isinstance(result_item, dict):
        return result_item.get("res", result_item)
    return {}


def process_layout_result(result_item, doc, all_text, all_scores, page_label=None):
    """Walk one PPStructureV3 result (one page/image) and emit either
    Word tables or paragraphs into the docx + plain-text buffer."""

    if page_label:
        doc.add_heading(page_label, level=2)
        all_text.append(f"\n===== {page_label} =====\n")

    res = get_res_dict(result_item)

    table_html_list = [
        t.get("pred_html", "") for t in res.get("table_res_list", []) or []
    ]
    table_index = 0

    blocks = res.get("parsing_res_list", []) or []

    if blocks:
        for block in blocks:
            label = str(block.get("block_label", "text")).lower()
            content = block.get("block_content", "") or ""

            if "table" in label:
                if table_index < len(table_html_list):
                    grid = html_table_to_grid(table_html_list[table_index])
                    grid = correct_grid(grid)
                    add_docx_table(doc, grid)

                    all_text.append("[TABLE]")
                    for row in grid:
                        all_text.append(" | ".join(row))
                    all_text.append("[/TABLE]")

                    table_index += 1
                else:
                    # fallback: no matching html, dump raw content
                    doc.add_paragraph(content)
                    all_text.append(content)

            elif "image" in label or "figure" in label or "chart" in label:
                doc.add_paragraph("[Figure/Image region detected — not extracted]")
                all_text.append("[FIGURE]")

            elif content.strip():
                doc.add_paragraph(content)
                all_text.append(content)

    else:
        # fallback: no parsing_res_list available, just render any tables
        # found and dump the plain-OCR text if present
        for html in table_html_list:
            grid = correct_grid(html_table_to_grid(html))
            add_docx_table(doc, grid)
            all_text.append("[TABLE]")
            for row in grid:
                all_text.append(" | ".join(row))
            all_text.append("[/TABLE]")

        for ocr_res in res.get("overall_ocr_res", {}).get("rec_texts", []) or []:
            doc.add_paragraph(ocr_res)
            all_text.append(ocr_res)

    scores = res.get("overall_ocr_res", {}).get("rec_scores", []) or []
    all_scores.extend(scores)


# ==========================================================
# IMAGE / PDF DRIVERS
# ==========================================================

def run_structure(image_bgr):
    """Run PPStructureV3 on a single BGR image, return its single result item."""
    results = structure_engine.predict(image_bgr)
    results = list(results)  # predict() may return a generator
    return results[0] if results else {}


def process_image(file_path, doc, all_text, all_scores):
    image = cv2.imread(file_path)
    if image is None:
        raise Exception("Unable to load image.")

    image = preprocess_image(image)
    result = run_structure(image)
    process_layout_result(result, doc, all_text, all_scores)


def process_pdf(file_path, doc, all_text, all_scores):
    info = pdfinfo_from_path(file_path)
    total_pages = int(info["Pages"])

    print(f"\nTotal Pages: {total_pages}")

    for page_num in range(1, total_pages + 1):
        print(f"[{page_num}/{total_pages}] Processing...")

        page = convert_from_path(
            file_path, first_page=page_num, last_page=page_num
        )[0]

        image = cv2.cvtColor(np.array(page), cv2.COLOR_RGB2BGR)
        image = preprocess_image(image)

        result = run_structure(image)
        process_layout_result(
            result, doc, all_text, all_scores,
            page_label=f"PAGE {page_num}"
        )


# ==========================================================
# SAVE OUTPUTS
# ==========================================================

def average_confidence(scores):
    if not scores:
        return 0.0
    return round(sum(scores) / len(scores) * 100, 2)


def save_txt(output_folder, base_name, content):
    txt_file = os.path.join(output_folder, f"{base_name}.txt")
    with open(txt_file, "w", encoding="utf-8") as f:
        f.write("\n".join(content))
    return txt_file


# ==========================================================
# MAIN
# ==========================================================

def main():
    file_path = input("\nEnter image or PDF path: ").strip()

    if not os.path.exists(file_path):
        print("❌ File not found.")
        return

    start_time = time.time()

    base_name = os.path.splitext(os.path.basename(file_path))[0]
    output_folder = os.path.join(OUTPUT_DIR, f"{base_name}_output")
    os.makedirs(output_folder, exist_ok=True)

    doc = Document()
    doc.add_heading(f"OCR Output - {base_name}", level=1)

    all_text = []
    all_scores = []

    try:
        if file_path.lower().endswith(".pdf"):
            process_pdf(file_path, doc, all_text, all_scores)
        else:
            process_image(file_path, doc, all_text, all_scores)

        confidence = average_confidence(all_scores)

        doc.add_heading("OCR Statistics", level=2)
        stats_p = doc.add_paragraph(f"Average Confidence: {confidence}%")
        stats_p.runs[0].bold = True

        all_text.append("")
        all_text.append("=" * 40)
        all_text.append("OCR STATISTICS")
        all_text.append("=" * 40)
        all_text.append(f"Average Confidence: {confidence}%")

        txt_file = save_txt(output_folder, base_name, all_text)

        docx_file = os.path.join(output_folder, f"{base_name}.docx")
        doc.save(docx_file)

        elapsed = round(time.time() - start_time, 2)

        print("\n✓ OCR Completed")
        print(f"Confidence : {confidence}%")
        print(f"Time Taken : {elapsed}s")

        print("\nTXT:")
        print(txt_file)

        print("\nDOCX:")
        print(docx_file)

    except Exception as e:
        print(f"\n❌ Error: {e}")


if __name__ == "__main__":
    main()