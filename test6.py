from paddleocr import PaddleOCR
from pdf2image import convert_from_path
from docx import Document
import os

# -------------------------------
# Initialize OCR Model
# -------------------------------
print("Loading OCR model...")

ocr = PaddleOCR(
    lang="en",
    device="cpu",
    use_doc_orientation_classify=False,
    use_doc_unwarping=False,
    use_textline_orientation=False
)

# -------------------------------
# Input File
# -------------------------------
file_path = input("\nEnter image or PDF path: ").strip()

if not os.path.exists(file_path):
    print("Error: File not found!")
    exit()

# -------------------------------
# Create Output Folder
# -------------------------------
OUTPUT_DIR = "output"
os.makedirs(OUTPUT_DIR, exist_ok=True)

base_name = os.path.splitext(
    os.path.basename(file_path)
)[0]

output_folder = os.path.join(
    OUTPUT_DIR,
    f"{base_name}_output"
)

os.makedirs(output_folder, exist_ok=True)

all_text = []

# -------------------------------
# PDF Processing
# -------------------------------
if file_path.lower().endswith(".pdf"):

    print("\nConverting PDF to images...")

    pages = convert_from_path(file_path)

    total_pages = len(pages)

    print(f"Found {total_pages} page(s)")
    print("Running OCR...\n")

    for page_num, page in enumerate(pages, start=1):

        temp_image = os.path.join(
            output_folder,
            f"temp_page_{page_num}.png"
        )

        # PDF Page → Image
        page.save(temp_image, "PNG")

        # OCR
        result = ocr.predict(temp_image)

        all_text.append(
            f"\n===== PAGE {page_num} =====\n"
        )

        for page_result in result:
            for text in page_result["rec_texts"]:
                all_text.append(text)

        # Delete temporary image
        os.remove(temp_image)

        print(
            f"✓ Page {page_num}/{total_pages} processed"
        )

# -------------------------------
# Image Processing
# -------------------------------
else:

    print("\nRunning OCR on image...\n")

    result = ocr.predict(file_path)

    for page_result in result:
        for text in page_result["rec_texts"]:
            all_text.append(text)

    print("✓ Image processed")

# -------------------------------
# Save TXT File
# -------------------------------
print("\nSaving TXT file...")

txt_file = os.path.join(
    output_folder,
    f"{base_name}.txt"
)

with open(txt_file, "w", encoding="utf-8") as file:
    file.write("\n".join(all_text))

# -------------------------------
# Save DOCX File
# -------------------------------
print("Saving DOCX file...")

doc = Document()

doc.add_heading(
    f"OCR Output - {base_name}",
    level=1
)

for line in all_text:
    doc.add_paragraph(line)

docx_file = os.path.join(
    output_folder,
    f"{base_name}.docx"
)

doc.save(docx_file)

# -------------------------------
# Complete
# -------------------------------
print("\nOCR Completed Successfully!")
print(f"TXT  : {txt_file}")
print(f"DOCX : {docx_file}")