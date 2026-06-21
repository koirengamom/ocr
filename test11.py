import os
import cv2
import time
import warnings
import numpy as np

from paddleocr import PaddleOCR
from pdf2image import convert_from_path, pdfinfo_from_path
from docx import Document

# ==========================================================
# SETTINGS
# ==========================================================

warnings.filterwarnings("ignore")

os.environ["GLOG_minloglevel"] = "3"
os.environ["FLAGS_logtostderr"] = "0"
os.environ["FLAGS_use_mkldnn"] = "0"

OUTPUT_DIR = "output"

# ==========================================================
# IMAGE PREPROCESSING
# ==========================================================

def preprocess_image(image):

    gray = cv2.cvtColor(
        image,
        cv2.COLOR_BGR2GRAY
    )

    gray = cv2.fastNlMeansDenoising(
        gray,
        None,
        10,
        7,
        21
    )

    processed = cv2.cvtColor(
        gray,
        cv2.COLOR_GRAY2BGR
    )

    return processed


# ==========================================================
# OCR ENGINE
# ==========================================================

print("Loading OCR model...")

ocr = PaddleOCR(
    lang="en",
    device="cpu",
    enable_mkldnn=False,
    use_doc_orientation_classify=True,
    use_doc_unwarping=True,
    use_textline_orientation=True
)

print("✓ OCR model loaded")


# ==========================================================
# OCR HELPERS
# ==========================================================

def extract_text(result):

    texts = []
    confidences = []

    for page_result in result:

        if "rec_texts" in page_result:

            texts.extend(
                page_result["rec_texts"]
            )

        if "rec_scores" in page_result:

            confidences.extend(
                page_result["rec_scores"]
            )

    return texts, confidences


def average_confidence(scores):

    if not scores:
        return 0.0

    return round(
        sum(scores) / len(scores) * 100,
        2
    )


# ==========================================================
# IMAGE OCR
# ==========================================================

def process_image(file_path):

    image = cv2.imread(file_path)

    if image is None:
        raise Exception(
            "Unable to load image."
        )

    image = preprocess_image(image)

    result = ocr.predict(image)

    return extract_text(result)


# ==========================================================
# PDF OCR
# ==========================================================

def process_pdf(file_path):

    info = pdfinfo_from_path(file_path)

    total_pages = int(info["Pages"])

    print(
        f"\nTotal Pages: {total_pages}"
    )

    all_text = []
    all_scores = []

    for page_num in range(
        1,
        total_pages + 1
    ):

        print(
            f"[{page_num}/{total_pages}] "
            f"Processing..."
        )

        page = convert_from_path(
            file_path,
            first_page=page_num,
            last_page=page_num
        )[0]

        image = cv2.cvtColor(
            np.array(page),
            cv2.COLOR_RGB2BGR
        )

        image = preprocess_image(image)

        result = ocr.predict(image)

        texts, scores = extract_text(result)

        all_text.append(
            f"\n===== PAGE {page_num} =====\n"
        )

        all_text.extend(texts)

        all_scores.extend(scores)

    return all_text, all_scores


# ==========================================================
# SAVE OUTPUTS
# ==========================================================

def save_txt(
    output_folder,
    base_name,
    content
):

    txt_file = os.path.join(
        output_folder,
        f"{base_name}.txt"
    )

    with open(
        txt_file,
        "w",
        encoding="utf-8"
    ) as file:

        file.write(
            "\n".join(content)
        )

    return txt_file


def save_docx(
    output_folder,
    base_name,
    content
):

    doc = Document()

    doc.add_heading(
        f"OCR Output - {base_name}",
        level=1
    )

    for line in content:

        doc.add_paragraph(
            str(line)
        )

    docx_file = os.path.join(
        output_folder,
        f"{base_name}.docx"
    )

    doc.save(docx_file)

    return docx_file


# ==========================================================
# MAIN
# ==========================================================

def main():

    file_path = input(
        "\nEnter image or PDF path: "
    ).strip()

    if not os.path.exists(file_path):

        print(
            "❌ File not found."
        )
        return

    start_time = time.time()

    base_name = os.path.splitext(
        os.path.basename(file_path)
    )[0]

    output_folder = os.path.join(
        OUTPUT_DIR,
        f"{base_name}_output"
    )

    os.makedirs(
        output_folder,
        exist_ok=True
    )

    try:

        if file_path.lower().endswith(
            ".pdf"
        ):

            texts, scores = process_pdf(
                file_path
            )

        else:

            texts, scores = process_image(
                file_path
            )

        confidence = average_confidence(
            scores
        )

        texts.append("")
        texts.append("=" * 40)
        texts.append("OCR STATISTICS")
        texts.append("=" * 40)
        texts.append(
            f"Average Confidence: "
            f"{confidence}%"
        )

        txt_file = save_txt(
            output_folder,
            base_name,
            texts
        )

        docx_file = save_docx(
            output_folder,
            base_name,
            texts
        )

        elapsed = round(
            time.time() - start_time,
            2
        )

        print("\n✓ OCR Completed")
        print(
            f"Confidence : "
            f"{confidence}%"
        )
        print(
            f"Time Taken : "
            f"{elapsed}s"
        )

        print("\nTXT:")
        print(txt_file)

        print("\nDOCX:")
        print(docx_file)

    except Exception as e:

        print(
            f"\n❌ Error: {e}"
        )


if __name__ == "__main__":
    main()